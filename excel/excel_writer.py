import os
from config.config import Config
from utils.logger import setup_logger
import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

logger = setup_logger()

def create_report(data):
    # CREATE THE OUTPUT FOLDER
    os.makedirs(Config.OUTPUT_FOLDER, exist_ok=True)

    save_path = os.path.join(Config.OUTPUT_FOLDER, Config.OUTPUT_FILENAME)

    try:
        with pd.ExcelWriter(save_path) as writer:
            data.to_excel(writer, sheet_name='data_products', index=False)

            # FORMAT THE DATA PRODUCTS SHEET
            worksheet = writer.sheets['data_products']

            for cell in worksheet[1]:
                # TEXT FORMATTING IN HEADERS
                cell.font = Font(bold=True, size=12, color='FFFFFF')
                cell.alignment = Alignment(horizontal='center')
                cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")

                # FIND THE PRICE COLUMNN
                if cell.value == 'price':
                    price_column = cell.column

            # NUMBER FORMATTING IN PRICE COLUMN
            for row in worksheet.iter_rows(min_row=2, min_col=price_column,max_col=price_column):
                for cell in row:
                    cell.number_format = "#,##0.00"


            for column_number, column in enumerate(worksheet.columns, start=1):
                max_length = 0

                for cell in column:
                    if cell.value is not None:
                        length = len(str(cell.value))

                        max_length = max(max_length, length)

                # GET THE COLUMN LETTER
                column_letter = get_column_letter(column_number)

                # ADJUST THE COLUMN WIDTH
                worksheet.column_dimensions[column_letter].width = min(max_length + 2, 40)

            # FREEZE THE ROW 1 HEADERS
            worksheet.freeze_panes = 'A2'

            # ADD AUTOFILTER
            worksheet.auto_filter.ref = worksheet.dimensions

        logger.info(f'File successfuly saved {save_path}.')

    except PermissionError:
        print(f"Unable to save {Config.OUTPUT_FILENAME}. Please close the file and try again.")